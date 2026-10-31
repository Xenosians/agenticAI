"""Owner-reviewed PDF/DOCX reference extraction stays inside RAG admission."""
import asyncio
import base64
from io import BytesIO
import json
import subprocess
import sys
from types import SimpleNamespace

import pytest

pytest.importorskip("lupa.lua54")  # The document worker reuses the Lua workers' sandbox.

from docx import Document  # noqa: E402
from pypdf import PdfWriter
from pypdf.constants import UserAccessPermissions
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
)

from services.extensions import service
from services.extensions import document_ingestion
from services.extensions.document_ingestion import (
    DocumentIngestionError,
    MAX_BASE64_BYTES,
    extract_document,
)
from services.extensions.store import ExtensionStore
from services.extensions.transform import worker_environment


REFERENCE = {
    "name": "Operations manual",
    "source": "Owned lab manual",
    "license": "Owner-permitted reference use",
    "purpose": "rag",
    "reviewed": True,
}


def encode(data: bytes) -> str:
    return base64.b64encode(data).decode("ascii")


def docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Router maintenance requires a planned window.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Change"
    table.cell(0, 1).text = "Approved window"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def pdf_bytes(text: str = "Router maintenance requires a planned window.", password: str | None = None,
              permissions: int | None = None) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=180)
    font = writer._add_object(DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    }))
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
    })
    stream = DecodedStreamObject()
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream.set_data(f"BT /F1 12 Tf 20 100 Td ({escaped}) Tj ET".encode())
    page[NameObject("/Contents")] = writer._add_object(stream)
    if password is not None:
        options = {} if permissions is None else {"permissions_flag": UserAccessPermissions(permissions)}
        writer.encrypt(user_password=password, owner_password="owner-only", algorithm="AES-256", **options)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def cid_pdf_bytes() -> bytes:
    """A Type0/CID-font page whose ToUnicode map makes pypdf load a UTF-16 codec while parsing."""
    def stream(data: bytes) -> bytes:
        return b"<< /Length %d >>\nstream\n" % len(data) + data + b"\nendstream"
    cmap = (b"/CIDInit /ProcSet findresource begin 12 dict begin begincmap\n"
            b"/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def\n"
            b"/CMapName /Adobe-Identity-UCS def /CMapType 2 def\n"
            b"1 begincodespacerange <0000> <FFFF> endcodespacerange\n"
            b"3 beginbfchar <0001> <0056> <0002> <0050> <0003> <004E> endbfchar\n"
            b"endcmap CMapName currentdict /CMap defineresource pop end end")
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>",
               b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 180] /Contents 4 0 R "
               b"/Resources << /Font << /F1 5 0 R >> >> >>",
               stream(b"BT /F1 12 Tf 20 100 Td <000100020003> Tj ET"),
               b"<< /Type /Font /Subtype /Type0 /BaseFont /Fixture /Encoding /Identity-H "
               b"/DescendantFonts [6 0 R] /ToUnicode 7 0 R >>",
               b"<< /Type /Font /Subtype /CIDFontType2 /BaseFont /Fixture "
               b"/CIDSystemInfo << /Registry (Adobe) /Ordering (Identity) /Supplement 0 >> /DW 500 >>",
               stream(cmap)]
    out, offsets = bytearray(b"%PDF-1.7\n"), []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return bytes(out)


@pytest.mark.parametrize(
    ("filename", "data", "expected"),
    [
        ("manual.docx", docx_bytes(), "Approved window"),
        ("manual.pdf", pdf_bytes(), "planned window"),
    ],
)
def test_supported_reference_documents_extract_text_and_record_source_digest(filename, data, expected):
    extracted = asyncio.run(extract_document(filename, encode(data)))
    assert expected in extracted.text
    assert extracted.metadata["filename"] == filename
    assert extracted.metadata["file_format"] == filename.rsplit(".", 1)[1]
    assert extracted.metadata["file_sha256"]
    assert extracted.metadata["parser_version"]


def test_document_upload_flows_through_reviewed_owner_scoped_dataset_store(tmp_path, monkeypatch):
    store = ExtensionStore(tmp_path)
    monkeypatch.setattr(service, "extension_store", lambda: store)
    settings = SimpleNamespace(model_profiles={})
    row = asyncio.run(service.command("alice", "dataset", {
        **REFERENCE, "filename": "manual.docx", "base64": encode(docx_bytes()),
    }, settings))

    assert row["kind"] == "dataset"
    assert row["document"]["file_format"] == "docx"
    assert "content" not in row and "base64" not in row
    hits = store.search("alice", row["id"], "planned window")
    assert hits and hits[0]["dataset_id"] == row["id"]
    assert store.list("bob") == []


@pytest.mark.parametrize(
    ("filename", "data", "code"),
    [
        ("manual.docx", b"not a zip archive", "document_file_invalid"),
        ("manual.pdf", pdf_bytes(password="owner-password"), "document_encrypted"),
        ("manual.pdf", pdf_bytes(""), "document_content_empty"),
        ("manual.pdf", b"%PDF-1.7\nnot really a document", "document_file_invalid"),
    ],
)
def test_invalid_encrypted_and_empty_documents_are_refused(filename, data, code):
    with pytest.raises(DocumentIngestionError) as raised:
        asyncio.run(extract_document(filename, encode(data)))
    assert raised.value.code == code


def test_a_pdf_encrypted_only_to_declare_permissions_is_read():
    # Vendor manuals are often encrypted with an empty user password; anyone may open them.
    extracted = asyncio.run(extract_document("vendor.pdf", encode(pdf_bytes(password=""))))
    assert "planned window" in extracted.text


def test_a_pdf_whose_author_forbids_text_extraction_is_refused():
    everything_but_extract = 0xFFFFFFFC & ~UserAccessPermissions.EXTRACT
    data = pdf_bytes(password="", permissions=everything_but_extract)
    with pytest.raises(DocumentIngestionError) as raised:
        asyncio.run(extract_document("locked.pdf", encode(data)))
    assert raised.value.code == "document_extraction_not_permitted"


def test_a_font_whose_text_map_needs_a_codec_is_read_inside_the_sandbox():
    """Parsing this page loads a UTF-16 codec, which a locked worker could no longer open from disk."""
    extracted = asyncio.run(extract_document("cid.pdf", encode(cid_pdf_bytes())))
    assert extracted.text.strip() == "VPN"


def test_documents_are_parsed_in_the_locked_worker_with_no_credentials(monkeypatch):
    """The worker gets the Lua workers' sandbox and minimal environment, never the service's."""
    monkeypatch.setenv("ITSM_INTERNAL_JOB_TOKEN", "must-not-reach-the-parser")
    seen = {}
    real_exec = asyncio.create_subprocess_exec

    async def spy(*args, **kwargs):
        seen.update(args=args, env=kwargs["env"], close_fds=kwargs.get("close_fds"))
        return await real_exec(*args, **kwargs)

    monkeypatch.setattr(document_ingestion.asyncio, "create_subprocess_exec", spy)
    asyncio.run(extract_document("manual.pdf", encode(pdf_bytes())))
    assert seen["args"][1:3] == ("-I", "-B") and seen["close_fds"] is True
    assert "ITSM_INTERNAL_JOB_TOKEN" not in seen["env"]
    assert set(seen["env"]) <= {"PATH", "LANG", "PYTHONDONTWRITEBYTECODE", "PYTHONNOUSERSITE"}


def test_after_lockdown_the_document_worker_cannot_open_files_sockets_or_processes(tmp_path):
    """Run the worker's real preload and lockdown, then try each escape inside the locked process."""
    secret = tmp_path / "secret.txt"
    secret.write_text("private")
    probe = f"""
import json, runpy, socket, subprocess, sys
worker = runpy.run_path({json.dumps(str(document_ingestion.WORKER))}, run_name="probe")
worker["preload"]()
layers = worker["shared"]["lockdown"](5, 512 * 1024 * 1024)
results = {{"layers": layers}}
for name, attempt in (
    ("read", lambda: open({json.dumps(str(secret))}).read()),
    ("socket", lambda: socket.create_connection(("127.0.0.1", 4000), timeout=1)),
    ("process", lambda: subprocess.run(["/bin/true"], check=True)),
):
    try:
        attempt()
        results[name] = "allowed"
    except Exception as exc:
        results[name] = type(exc).__name__
print(json.dumps(results))
"""
    completed = subprocess.run([sys.executable, "-I", "-c", probe], capture_output=True, text=True, timeout=60,
                               env=worker_environment())
    results = json.loads(completed.stdout.strip().splitlines()[-1])
    assert document_ingestion.REQUIRED_ISOLATION <= set(results["layers"])
    for name in ("read", "socket", "process"):
        assert results[name] != "allowed", (name, results)


@pytest.mark.parametrize(("answer", "code"), [
    ({"error": "provider_secret_leak"}, "document_parse_failed"),
    ({"text": "Some text", "isolation": ["process"]}, "document_parser_unavailable"),
    ({"text": 7, "isolation": ["process"]}, "document_parse_failed"),
])
def test_a_malformed_or_weakly_isolated_worker_answer_is_never_trusted(answer, code):
    with pytest.raises(DocumentIngestionError) as raised:
        document_ingestion._worker_text(0, json.dumps(answer).encode())
    assert raised.value.code == code


def test_a_worker_killed_by_its_limits_is_reported_as_a_resource_limit():
    with pytest.raises(DocumentIngestionError) as raised:
        document_ingestion._worker_text(-9, b"")
    assert raised.value.code == "document_resource_limit"


def test_document_payload_cannot_enter_training_admission(tmp_path, monkeypatch):
    store = ExtensionStore(tmp_path)
    monkeypatch.setattr(service, "extension_store", lambda: store)
    settings = SimpleNamespace(model_profiles={})
    payload = {
        **REFERENCE, "purpose": "training", "filename": "manual.docx",
        "base64": encode(docx_bytes()),
    }
    with pytest.raises(DocumentIngestionError) as raised:
        asyncio.run(service.command("alice", "dataset", payload, settings))
    assert raised.value.code == "document_training_not_supported"
    assert store.list("alice") == []


def test_filename_paths_and_oversized_encoded_payload_are_refused():
    with pytest.raises(DocumentIngestionError) as raised:
        asyncio.run(extract_document("../manual.pdf", encode(pdf_bytes())))
    assert raised.value.code == "document_file_invalid"
    with pytest.raises(DocumentIngestionError) as raised:
        asyncio.run(extract_document("manual.pdf", "A" * 2_666_669))
    assert raised.value.code == "document_upload_too_large"


def test_oversized_document_envelope_is_rejected_before_lua_admission(monkeypatch):
    from services.lua_adapters import host_boundaries

    inspected = False

    async def inspect(*args, **kwargs):
        nonlocal inspected
        inspected = True

    monkeypatch.setattr(host_boundaries, "inspect_data_async", inspect)
    with pytest.raises(DocumentIngestionError) as raised:
        asyncio.run(service.command("alice", "dataset", {
            **REFERENCE, "filename": "manual.pdf", "base64": "A" * (MAX_BASE64_BYTES + 1),
        }, SimpleNamespace(model_profiles={})))
    assert raised.value.code == "document_upload_too_large"
    assert not inspected


def test_internal_api_returns_only_a_curated_document_refusal(tmp_path, monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from api import extensions

    monkeypatch.setenv("EXTENSION_STORAGE_ROOT", str(tmp_path))
    settings = SimpleNamespace(
        require_internal_job_token=lambda: "document-test-token",
        model_profiles={},
    )
    monkeypatch.setattr(extensions, "get_settings", lambda: settings)
    service.extension_store.cache_clear()
    app = FastAPI()
    app.include_router(extensions.router)
    client = TestClient(app)
    response = client.post(
        "/internal/v1/extensions",
        headers={"x-internal-token": "document-test-token"},
        json={"owner": "alice", "action": "dataset", "payload": {
            **REFERENCE, "filename": "manual.docx", "base64": encode(b"not a document"),
        }},
    )
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "document_file_invalid"
    assert "not a document" not in response.text
    assert service.extension_store().list("alice") == []
    service.extension_store.cache_clear()
